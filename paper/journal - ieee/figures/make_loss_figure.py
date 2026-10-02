"""Figure for Section 6.1: why LoCa, DKD and RLD did not surpass classical KD (student CNN-S).

(a) Median ratio of the weighted response term to the weighted cross-entropy against the mean
    accuracy change from classical KD, with 95% bootstrap intervals across the 52 settings.
(b) Target-recall change of the published DKD and RLD from classical KD on each dataset, averaged
    over its four imbalance factors, with datasets ordered by class count. Spearman's rho between
    class count and that change is computed over the 13 datasets.
(c) LoCa's accuracy change from classical KD against the teacher's error on its own augmented
    training images in its final training epoch, one point per setting, with Spearman's r_s. Larger
    filled diamonds give the mean change within error bands; open markers are used only in (a),
    for the lower-weight DKD and RLD.

Run from the repository root:  python paper/journal/figures/make_loss_figure.py
"""

import json
import random
import statistics
from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt  # noqa: E402
from openpyxl import load_workbook  # noqa: E402
from scipy.stats import spearmanr  # noqa: E402

ROOT = Path(__file__).resolve().parents[3]
CATALOG = ROOT / "kd/catalog/v1.json"
NAMES = {"dermamnist": "DermaMNIST", "bloodmnist": "BloodMNIST", "pathmnist": "PathMNIST",
         "cifar10": "CIFAR-10", "cinic10": "CINIC-10", "eurosat": "EuroSAT",
         "fashionmnist": "Fashion-MNIST", "stl10": "STL-10", "svhn": "SVHN",
         "organamnist": "OrganAMNIST", "gtsrb": "GTSRB", "cifar100": "CIFAR-100",
         "caltech101": "Caltech-101"}
RUNS = ROOT / "runs"
MATRIX = RUNS / "experiment1-multidataset"
WORKBOOK = ROOT / "compiled-results/full-matrix-20260928/full-matrix-results-20260928.xlsx"
OUT = Path(__file__).with_name("loss_diagnostics.pdf")
SEEDS = (42, 43, 44)
WARMUP = 5
TEACHERS = {"CNN-T": ("cnn-cnn", None), "ResNet-18": ("r18-cnn", "r18"), "ResNet-34": ("r34-cnn", "r34")}
# Okabe-Ito colours, readable in colour-blind vision and in grey.
COLOURS = {"CNN-T": "#0072B2", "ResNet-18": "#D55E00", "ResNet-34": "#009E73"}
METHODS = {  # key: (workbook label, legend label, marker, filled)
    "kd": ("Classical KD", "Classical KD", "o", True),
    "loca": ("LoCa", "LoCa", "D", True),
    "dkd": ("DKD (published)", "DKD", "s", True),
    "dkd_matched": ("DKD (lower weight)", "DKD, lower weight", "s", False),
    "rld": ("RLD (published)", "RLD", "^", True),
    "rld_matched": ("RLD (lower weight)", "RLD, lower weight", "^", False),
}


def read(path):
    return json.loads(Path(path).read_text())


def jobs():
    return read(MATRIX / "matrix_summary.json")["jobs"]


def run_dir(teacher, method, job, seed):
    pairing, short = TEACHERS[teacher]
    if method == "kd":
        if short is None:
            return MATRIX / job["dataset"] / job["profile"] / "baselines/confirmatory" / f"kd_{job['id']}_seed{seed}"
        name = f"{job['id']}-{short}"
        return RUNS / "teacher-architecture" / name / "baselines/confirmatory" / f"kd_{name}_seed{seed}"
    return RUNS / "alternative-losses" / pairing / job["id"] / f"{method}_{job['id']}_seed{seed}"


def run_balance(directory):
    """Median over post-warm-up epochs of weighted response / weighted cross-entropy (as trained)."""
    config, history = read(directory / "config.json"), read(directory / "history.json")
    d = config["distillation"]
    ce_weight = 1.0 - d["weight"] if d["method"] in {"kd", "loca"} else 1.0
    return statistics.median(d["weight"] * e["train"]["response"] / (ce_weight * e["train"]["ce"])
                             for e in history[WARMUP:] if e["train"]["ce"] > 0)


def setting_scores():
    rows = list(load_workbook(WORKBOOK, read_only=True)["Losses by setting"].iter_rows(values_only=True))
    header = rows[0]
    table = [dict(zip(header, r)) for r in rows[1:] if r[header.index("Student")] == "CNN-S"]
    out = {}
    for teacher in TEACHERS:
        mine = [r for r in table if r["Teacher"] == teacher]
        assert len(mine) == 52, (teacher, len(mine))
        out[teacher] = {m: [(r[f"{lab} accuracy (%)"], r[f"{lab} target recall (%)"]) for r in mine]
                        for m, (lab, *_rest) in METHODS.items()}
    return out


def bootstrap(values, reps=2000, seed=2026):
    rng = random.Random(seed)
    means = sorted(statistics.fmean(rng.choices(values, k=len(values))) for _ in range(reps))
    return means[int(0.025 * reps)], means[int(0.975 * reps) - 1]


def main():
    scores = setting_scores()
    balance, change, teacher_acc = {}, {}, {}
    for teacher in TEACHERS:
        for method in METHODS:
            ratios = [run_balance(run_dir(teacher, method, j, s)) for j in jobs() for s in SEEDS]
            balance[(teacher, method)] = statistics.median(ratios)
            kd = scores[teacher]["kd"]
            acc = [m[0] - k[0] for m, k in zip(scores[teacher][method], kd)]
            rec = [m[1] - k[1] for m, k in zip(scores[teacher][method], kd)]
            change[(teacher, method)] = (statistics.fmean(acc), statistics.fmean(rec), bootstrap(acc))
        short = TEACHERS[teacher][1]
        teacher_acc[teacher] = []
        for j in jobs():
            if short is None:
                h = MATRIX / j["dataset"] / j["profile"] / "baselines/confirmatory" / f"teacher_{j['id']}" / "history.json"
            else:
                name = f"{j['id']}-{short}"
                h = RUNS / "teacher-architecture" / name / "baselines/confirmatory" / f"teacher_{name}" / "history.json"
            teacher_acc[teacher].append(read(h)[-1]["train"]["accuracy"] * 100)

    for (teacher, method), ratio in balance.items():
        a, r, (lo, hi) = change[(teacher, method)]
        print(f"{teacher:9s} {method:12s} ratio {ratio:6.2f}  dAcc {a:+.2f} [{lo:+.2f}, {hi:+.2f}]  dRtar {r:+.2f}")
    for teacher, values in teacher_acc.items():
        print(f"{teacher:9s} teacher train accuracy median {statistics.median(values):.1f}%  min {min(values):.1f}%")

    # (b) recall change of DKD and RLD on each dataset; (c) LoCa against teacher training error
    catalog = json.loads(CATALOG.read_text())["datasets"]
    names = [j["dataset"] for j in jobs()]
    datasets = sorted(set(names), key=lambda d: (catalog[d]["classes"], d))
    by_dataset, loca = {}, {}
    for teacher in TEACHERS:
        kd = scores[teacher]["kd"]
        for method in ("dkd", "rld", "dkd_matched", "rld_matched"):
            rec = [m[1] - k[1] for m, k in zip(scores[teacher][method], kd)]
            means = [statistics.fmean(r for r, n in zip(rec, names) if n == d) for d in datasets]
            by_dataset[(teacher, method)] = means
            rho, p = spearmanr([catalog[d]["classes"] for d in datasets], means)
            print(f"{teacher:9s} {method:12s} dRtar by dataset: " + " ".join(f"{NAMES[d]} {v:+.1f}" for d, v in zip(datasets, means))
                  + f"  | rho(classes) {rho:+.2f} p={p:.3f}")
        error = [100 - a for a in teacher_acc[teacher]]
        for method in ("loca", "rld"):
            for k, lab in ((0, "dAcc"), (1, "dRtar")):
                d = [m[k] - kk[k] for m, kk in zip(scores[teacher][method], kd)]
                rho, p = spearmanr(error, d)
                print(f"{teacher:9s} {method:4s} {lab} vs teacher training error: rho {rho:+.2f} p={p:.3f}")
        loca[teacher] = (error, [m[0] - k[0] for m, k in zip(scores[teacher]["loca"], kd)],
                         spearmanr(error, [m[0] - k[0] for m, k in zip(scores[teacher]["loca"], kd)])[0])

    plt.rcParams.update({"font.family": "STIXGeneral", "mathtext.fontset": "stix", "font.size": 7,
                         "axes.titlesize": 7.5, "axes.labelsize": 7, "legend.fontsize": 6.2,
                         "xtick.labelsize": 6.5, "ytick.labelsize": 6.5, "axes.linewidth": 0.6,
                         "pdf.fonttype": 42, "ps.fonttype": 42})
    fig, (ax_a, ax_b, ax_c) = plt.subplots(1, 3, figsize=(7.16, 2.95),
                                           gridspec_kw={"width_ratios": [1.15, 0.95, 1.0]})

    # (a) loss balance against accuracy change
    ax_a.axhline(0, color="0.6", lw=0.6, ls="--", zorder=0)
    for teacher, colour in COLOURS.items():
        for published, lower in (("dkd", "dkd_matched"), ("rld", "rld_matched")):
            x0, x1 = balance[(teacher, published)], balance[(teacher, lower)]
            y0, y1 = change[(teacher, published)][0], change[(teacher, lower)][0]
            ax_a.annotate("", xy=(x1, y1), xytext=(x0, y0),
                          arrowprops={"arrowstyle": "->", "color": colour, "lw": 0.6, "alpha": 0.55,
                                      "shrinkA": 3.5, "shrinkB": 3.5})
        for method, (_lab, _leg, marker, filled) in METHODS.items():
            x = balance[(teacher, method)]
            y, _r, (lo, hi) = change[(teacher, method)]
            ax_a.errorbar(x, y, yerr=[[y - lo], [hi - y]], fmt="none", ecolor=colour, elinewidth=0.6,
                          capsize=1.5, alpha=0.8, zorder=2)
            ax_a.scatter(x, y, marker=marker, s=22, facecolors=colour if filled else "white",
                         edgecolors=colour, linewidths=0.9, zorder=3)
    ax_a.set_xscale("log")
    ax_a.set_xticks([3, 5, 10, 20, 30])
    ax_a.get_xaxis().set_major_formatter(matplotlib.ticker.ScalarFormatter())
    ax_a.set_xlim(2.8, 38)
    ax_a.set_xlabel("Response term / cross-entropy (median)")
    ax_a.set_ylabel("Accuracy change from classical KD (pp)")
    ax_a.set_title("(a) Loss balance")

    # (b) recall change of the published DKD and RLD on each dataset, ordered by class count
    ax_b.axhline(0, color="0.6", lw=0.6, ls="--", zorder=0)
    for t, (teacher, colour) in enumerate(COLOURS.items()):
        for k, method in enumerate(("dkd", "rld")):
            marker = METHODS[method][2]
            dx = (t - 1) * 0.24 + (k - 0.5) * 0.1
            ax_b.scatter([i + dx for i in range(len(datasets))], by_dataset[(teacher, method)],
                         marker=marker, s=10, color=colour, lw=0, zorder=3)
    for i in range(len(datasets) - 1):  # separate the class counts 7, 8, 9, 10, 11, 43, 100 and 101
        if catalog[datasets[i]]["classes"] != catalog[datasets[i + 1]]["classes"]:
            ax_b.axvline(i + 0.5, color="0.88", lw=0.5, zorder=0)
    ax_b.set_xticks(range(len(datasets)))
    ax_b.set_xticklabels([f"{NAMES[d]} ({catalog[d]['classes']})" for d in datasets], rotation=90, fontsize=5.6)
    ax_b.set_xlim(-0.6, len(datasets) - 0.4)
    ax_b.set_ylabel("Target-recall change from classical KD (pp)")
    ax_b.set_title("(b) DKD and RLD by dataset")

    # (c) LoCa's accuracy change against the teacher's error on its own training images
    ax_c.axhline(0, color="0.6", lw=0.6, ls="--", zorder=0)
    for i, (teacher, colour) in enumerate(COLOURS.items()):
        error, gain, rho = loca[teacher]
        ax_c.scatter(error, gain, marker="D", s=7, color=colour, alpha=0.35, lw=0, zorder=2)
        # mean change within error bands of <=0.1%, 0.1-1% and >1%, placed at each band's median error
        bands = [[(e, g) for e, g in zip(error, gain) if lo < e <= hi or (lo == 0 and e <= hi)]
                 for lo, hi in ((0, 0.1), (0.1, 1), (1, 100))]
        xs = [statistics.median(e for e, _g in b) for b in bands if b]
        ys = [statistics.fmean(g for _e, g in b) for b in bands if b]
        ax_c.plot(xs, ys, color=colour, lw=0.9, zorder=3)
        ax_c.scatter(xs, ys, marker="D", s=24, facecolors=colour, edgecolors="0.1", linewidths=0.7, zorder=4)
        ax_c.text(0.03, 0.97 - 0.075 * i, f"{teacher}: $r_s$ = {rho:+.2f}".replace("-", "\u2212"),
                  transform=ax_c.transAxes, color=colour, fontsize=6, ha="left", va="top")
    ax_c.set_xscale("symlog", linthresh=0.1, linscale=0.6)
    ax_c.set_xticks([0, 0.1, 1, 10, 40])
    ax_c.set_xticklabels(["0", "0.1", "1", "10", "40"])
    ax_c.set_xlim(-0.02, 60)
    ax_c.set_xlabel("Teacher error on own training images (%)")
    ax_c.set_ylabel("LoCa accuracy change from classical KD (pp)")
    ax_c.set_title("(c) LoCa and teacher fit")

    # shared legends: methods (shape) and teachers (colour)
    from matplotlib.lines import Line2D
    method_handles = [Line2D([], [], ls="none", marker=m, markersize=4.2, markerfacecolor="0.25" if f else "white",
                             markeredgecolor="0.25", markeredgewidth=0.9, label=leg)
                      for _lab, leg, m, f in METHODS.values()]
    teacher_handles = [Line2D([], [], ls="none", marker="o", markersize=4.2, color=c, label=t)
                       for t, c in COLOURS.items()]
    fig.legend(handles=method_handles + teacher_handles, loc="lower center", ncol=9, frameon=False,
               bbox_to_anchor=(0.5, -0.015), handletextpad=0.25, columnspacing=0.9)
    for ax in (ax_a, ax_b, ax_c):
        ax.tick_params(width=0.6, length=2.5)
        for side in ("top", "right"):
            ax.spines[side].set_visible(False)
    fig.tight_layout(rect=(0, 0.075, 1, 1), w_pad=1.4)
    fig.savefig(OUT, bbox_inches="tight", pad_inches=0.02)
    print("wrote", OUT)


if __name__ == "__main__":
    main()
